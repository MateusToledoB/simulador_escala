"""
API do simulador de escala: recebe as linhas da PEC (cargo, escala, quantidade) e devolve o
resultado de main.calcular em JSON, com os mesmos textos do relatório do terminal.

    uvicorn api:app --reload --port 8000
"""
import tempfile
import threading
from collections import Counter, defaultdict
from contextlib import contextmanager
from io import BytesIO
from pathlib import Path
from typing import Literal

from fastapi import FastAPI, HTTPException, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field

import relatorio
from core.settings import settings
from main import (ESTRATEGIAS, cenarios, frase_extras, horas_pessoa, minimo_de_pessoas, contrato, desc_turno, faixas, fmt_dias, fmt_hora, fmt_horas, fmt_min, frase_turno,
                  frases_horistas, horas_legais, horas_semana, horas_turno, intervalo, janela, ler_escala, ler_linhas, meta_semana,
                  montar_equipes, nome_turno, por_dia, sugerir, total_revezamento)

app = FastAPI(title="Simulador de escala")
app.add_middleware(CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])


class LinhaPEC(BaseModel):
    cargo: str
    escala: str = Field(description="como na PEC: 'ESCALA SEGUNDA - SEXTA (07:30 - 17:18, INTERVALO 12:00 - 13:00)', "
                                    "'ESCALA 6 X 1 (07:30 - 16:00)' ou dinâmica: "
                                    "'ESCALA SEGUNDA - QUINTA (08:00 - 19:00) / SEXTA (08:00 - 18:00)'; sem intervalo, "
                                    "usa o mínimo da CLT")
    quantidade: int = Field(gt=0)
    salario: float | None = Field(None, ge=0, description="salário médio mensal do cargo; custo = salário x FATOR_ENCARGOS")


class Simulacao(BaseModel):
    linhas: list[LinhaPEC] = Field(min_length=1)
    limite_horas: Literal[40, 42] = Field(42, description="carga semanal máxima de cada pessoa")
    salario: float | None = Field(None, ge=0, description="salário padrão das linhas sem salário próprio")

    def linhas_com_salario(self):
        return [{**l.model_dump(), "salario": l.salario or self.salario or None} for l in self.linhas]


_calculo = threading.Lock()
_salario = None  # salário das contratações em cálculo (_parametros / _salario_medio)


@contextmanager
def _parametros(simulacao):
    """
    Calcula com a carga semanal e o salário escolhidos; o lock impede que duas simulações com parâmetros
    diferentes se misturem.
    """
    global _salario
    with _calculo:
        antes = settings.LIMITE_HORAS_SEMANA
        settings.LIMITE_HORAS_SEMANA, _salario = simulacao.limite_horas, simulacao.salario or None
        try:
            yield
        finally:
            settings.LIMITE_HORAS_SEMANA, _salario = antes, None


def _salario_medio(pesos):
    """Média dos salários [(peso, salário)] ponderada pelo peso, só com as linhas que têm salário; None sem nenhum."""
    com = [(p, s) for p, s in pesos if s]
    if not com:
        return None
    total = sum(p for p, _ in com)
    return sum(p * s for p, s in com) / total if total > 0 else sum(s for _, s in com) / len(com)


def _perda(e):
    """Horas de posto por semana que a equipe deixa de cobrir com a nova escala (peso do salário das contratações)."""
    antes = sum(e["dias_antes"][d] * (janela(e["antes"][d])[1] - janela(e["antes"][d])[0]) for d in e["dias"])
    depois = sum(q * (janela(e["depois"][d])[1] - janela(e["depois"][d])[0]) for d, q in e["dias_depois"].items())
    return max(0, antes - depois) / 60


def _por_cargo(itens, sugestoes):
    """
    Impacto por cargo. itens = [(cargo, pessoas, salário, peso)], com peso = horas por semana que o cargo deixa de
    cobrir (quadro) ou de posto (demanda). O custo das contratações é repartido pelo peso de cada cargo, cada um
    pelo próprio salário: a soma dos cargos com salário dá o custo total das contratações.
    """
    g = {}
    for cargo, qtd, sal, peso in itens:
        o = g.setdefault(cargo, {"pessoas": 0, "folha": 0.0, "com_salario": 0, "peso": 0.0})
        o["pessoas"] += qtd
        o["peso"] += peso
        if sal:
            o["folha"] += sal * qtd
            o["com_salario"] += qtd
    total = sum(o["peso"] for o in g.values())
    # média dos salários pesada pelo peso, como em _salario_medio: base para repartir o custo das contratações
    pesos = [(o["peso"], o["folha"] / o["com_salario"]) for o in g.values() if o["com_salario"] and o["peso"] > 0]
    base = sum(p * s for p, s in pesos)
    custos = {k: (s or {}).get("custo_mensal") for k, s in (sugestoes or {}).items()}
    out = []
    for cargo, o in g.items():
        sal = o["folha"] / o["com_salario"] if o["com_salario"] else None
        fatia = o["peso"] * sal / base if sal and base else 0
        out.append({
            "cargo": cargo,
            "pessoas": o["pessoas"],
            "salario": None if sal is None else round(sal, 2),
            "custo_atual": round(o["folha"] * settings.FATOR_ENCARGOS, 2) if o["folha"] else None,
            "horas": fmt_horas(o["peso"]),
            "parte": o["peso"] / total if total else 0,
            **{f"impacto_{k}": None if c is None or sal is None else round(c * fatia, 2) for k, c in custos.items()},
        })
    return sorted(out, key=lambda x: (-x["parte"], -x["pessoas"]))


def _custo_quadro(linhas):
    """Custo mensal do quadro informado (quantidade x salário x FATOR_ENCARGOS das linhas com salário)."""
    com = [l for l in linhas if l.get("salario")]
    if not com:
        return None
    return {
        "mensal": round(sum(l["quantidade"] * l["salario"] for l in com) * settings.FATOR_ENCARGOS, 2),
        "pessoas_com_salario": sum(l["quantidade"] for l in com),
        "salario_medio": round(_salario_medio([(l["quantidade"], l["salario"]) for l in com]), 2),
    }


def _custo(pessoas=0, horas_horista=0.0):
    """
    Custo mensal (salário x FATOR_ENCARGOS): cada pessoa 5x2 custa um salário; o horista, a parte do salário
    proporcional às horas (horas na semana / LIMITE_HORAS_SEMANA). None sem salário informado.
    """
    if _salario is None:
        return None
    return round(_salario * settings.FATOR_ENCARGOS * (pessoas + horas_horista / settings.LIMITE_HORAS_SEMANA), 2)


# ---------------------------------------------------------------------------
# SERIALIZAÇÃO
# ---------------------------------------------------------------------------
def _turno_json(segs):
    ini, fim = janela(segs)
    return {
        "entrada": fmt_hora(ini),
        "saida": fmt_hora(fim),
        "intervalo": fmt_horas(intervalo(segs) / 60),
        "intervalo_inicio": fmt_hora(segs[0][1]) if len(segs) > 1 else None,
        "intervalo_fim": fmt_hora(segs[1][0]) if len(segs) > 1 else None,
        "horas": fmt_horas(horas_legais(segs)),
        "descricao": desc_turno(segs),
    }


def _ajuste(e):
    """O que mudou no turno: 'intervalo +24 min', 'saída 30 min mais cedo'... (vazio = nada)."""
    partes = []
    mais_intervalo = max(intervalo(e["depois"][d]) - intervalo(e["antes"][d]) for d in e["dias"])
    if mais_intervalo > 0:
        partes.append(f"intervalo +{fmt_horas(mais_intervalo / 60)}")
    if e["corte"] < 0:
        partes.append(f"saída {fmt_horas(-e['corte'] / 60)} mais tarde")
    elif e["corte"] > 0:
        quando = "entrada" if e["estrategia"] == "entrada" else "saída"
        sentido = "mais tarde" if e["estrategia"] == "entrada" else "mais cedo"
        partes.append(f"{quando} {fmt_horas(e['corte'] / 60)} {sentido}")
    return ", ".join(partes)


def _cargo_json(e):
    dias = len(e["pessoas"][0])
    return {
        "cargo": e["cargo"],
        "escala": e["escala"],
        "quantidade": e["quantidade"],
        "dias": fmt_dias(e["dias"]),
        "faixas": [{"dias": fmt_dias(d), "antes": _turno_json(antes), "depois": _turno_json(depois)}
                   for d, antes, depois in faixas(e)],
        "semana_antes": fmt_horas(horas_semana(e["antes"], [e["dias"]])),
        "semana_depois": fmt_horas(horas_semana(e["depois"], e["pessoas"])),
        "corte_min": e["corte"],
        "ajuste": _ajuste(e),
        "enquadrou": e["enquadrou"],
        "meta": fmt_horas(meta_semana(dias)),
        "rodizio": [{"quantidade": qtd, "dias": fmt_dias(list(d))}
                    for d, qtd in Counter(map(tuple, e["pessoas"])).items()] if dias < len(e["dias"]) else [],
    }


def _horistas_json(pessoas):
    ordem = sorted(pessoas, key=lambda p: -len(p))
    return {
        "total": len(pessoas),
        "custo_mensal": _custo(horas_horista=sum(horas_turno(*t) for p in pessoas for t in p.values())),
        "horas_semana": fmt_horas(sum(horas_turno(*t) for p in pessoas for t in p.values())),
        "frases": frases_horistas(pessoas),
        "pessoas": [{
            "contrato": contrato(p),
            "horas_semana": fmt_horas(sum(horas_turno(*t) for t in p.values())),
            "turnos": [{"dia": settings.NOMES[dia % 7], "entrada": fmt_hora(ini), "saida": fmt_hora(fim),
                        "horas": fmt_horas(horas_turno(ini, fim))} for dia, (ini, fim) in sorted(p.items())],
        } for p in ordem],
    }


def _revezamento_json(rev):
    """Equipes 5x2 por turno (com a grade semana x pessoa) e os horistas que cobrem o resto."""
    horas = sum(horas_pessoa(p) for g in rev["turnos"] for p in g["pessoas"])
    horas_horistas = sum(horas_turno(*t) for p in rev["horistas"] for t in p.values())
    horas += horas_horistas
    return {
        "total": total_revezamento(rev),
        "custo_mensal": _custo(sum(len(g["pessoas"]) for g in rev["turnos"]), horas_horistas),
        "horas_semana": fmt_horas(horas),
        "frases": [frase_turno(g) for g in rev["turnos"]] + frases_horistas(rev["horistas"]),
        "turnos": [{
            "titulo": f"Turno da {nome_turno(g['turno'][0])}",
            "entrada": fmt_hora(g["turno"][0]),
            "saida": fmt_hora(g["turno"][1]),
            "intervalo": f"{fmt_hora(g['intervalo'][0])}–{fmt_hora(g['intervalo'][1])}" if "intervalo" in g else None,
            "horas_dia": fmt_horas(horas_turno(*g["turno"])),
            "horas_semana": fmt_horas(horas_turno(*g["turno"]) * settings.DIAS_TRABALHO),
            "por_dia": por_dia(g),
            "demanda": g["falta"],
            "pessoas": [{"nome": chr(65 + i) if i < 26 else f"P{i + 1}", "folgas": p["folgas"],
                         "trabalha": [d in p["escala"] for d in range(7)],
                         "complemento": [d in p.get("extras", {}) for d in range(7)],
                         "extras": frase_extras(p) if p.get("extras") else None,
                         "horas_semana": fmt_horas(horas_pessoa(p))} for i, p in enumerate(g["pessoas"])],
        } for g in rev["turnos"]],
        "horistas": _horistas_json(rev["horistas"]),
    }


def _config_json():
    return {
        "limite_horas_semana": settings.LIMITE_HORAS_SEMANA,
        "salario": _salario,
        "fator_encargos": settings.FATOR_ENCARGOS,
        "escala": f"{settings.DIAS_TRABALHO}x{settings.FOLGAS_POR_SEMANA}",
        "intervalo_maximo": fmt_horas(settings.INTERVALO_MAXIMO / 60),
        "corte_maximo_min": settings.CORTE_MAXIMO,
        "dias": settings.NOMES,
    }


def _sugestoes_json(r):
    return {"revezamento_puro": _revezamento_json(r["revezamento_puro"]),
            "revezamento": _revezamento_json(r["revezamento"]),
            "clt_horista": _revezamento_json(r["clt_horista"])}


def _opcao_json(estrategia, r, fora):
    """
    As contratações custam o salário médio dos cargos que perderam cobertura, ponderado pelas horas perdidas.
    fora = linhas do quadro sem horário reconhecido: entram só no custo atual por cargo.
    """
    global _salario
    titulo, descricao = ESTRATEGIAS[estrategia]
    _salario = (_salario_medio([(_perda(e), e.get("salario")) for e in r["equipes"]])
                or _salario_medio([(e["quantidade"], e.get("salario")) for e in r["equipes"]]))
    sugestoes = _sugestoes_json(r) if r["blocos"] else None
    itens = ([(e["cargo"], e["quantidade"], e.get("salario"), _perda(e)) for e in r["equipes"]]
             + [(l["cargo"], l["quantidade"], l.get("salario"), 0) for l in fora])
    return {
        "id": estrategia,
        "titulo": titulo,
        "descricao": descricao,
        "salario_contratacao": None if _salario is None else round(_salario, 2),
        "cargos": [_cargo_json(e) for e in r["equipes"]],
        "rombo": {
            "blocos": [{"inicio": fmt_min(ini), "fim": fmt_min(fim), "duracao": fmt_horas((fim - ini) / 60),
                        "pessoas": qtd} for (ini, fim), qtd in sorted(Counter(r["blocos"]).items())],
            "por_dia": [{"dia": settings.NOMES[d], "horas": fmt_horas(h), "valor": round(h, 2)}
                        for d, h in r["horas_por_dia"].items()],
            "total_horas": fmt_horas(r["total_horas"]),
            "horas_demandadas": fmt_horas(r["horas_demandadas"]),
            "percentual": r["total_horas"] / r["horas_demandadas"] if r["horas_demandadas"] else 0,
            "colaboradores_equivalentes": round(r["total_horas"] / settings.LIMITE_HORAS_SEMANA, 1),
        },
        "sugestoes": sugestoes,
        "por_cargo": _por_cargo(itens, sugestoes),
    }


def resultado_json(opcoes, equipes, sem_horario, linhas):
    """opcoes = {estrategia: resultado de calcular} (main.cenarios); linhas = as do quadro, com o salário."""
    return {
        "config": _config_json(),
        "total_pessoas": sum(e["quantidade"] for e in equipes) + sum(s[3] for s in sem_horario),
        "custo_quadro": _custo_quadro(linhas),
        "sem_horario": [{"cargo": c, "escala": esc, "quantidade": qtd} for _, c, esc, qtd in sem_horario],
        "opcoes": [_opcao_json(estrategia, r, [l for l in linhas if not ler_escala(l["escala"])])
                   for estrategia, r in opcoes.items()],
    }


# ---------------------------------------------------------------------------
# ROTAS
# ---------------------------------------------------------------------------
@app.get("/health")
def health():
    return {"ok": True}


@app.post("/simular")
def simular(simulacao: Simulacao):
    linhas = simulacao.linhas_com_salario()
    equipes, sem_horario = montar_equipes(linhas)
    with _parametros(simulacao):
        return resultado_json(cenarios(equipes), equipes, sem_horario, linhas)


PROPOSTAS = {
    "horario": ("Manter o horário do posto",
                "Cobre o posto em todo o horário informado, inclusive no intervalo de cada pessoa."),
    "minimo": ("Mínimo de colaboradores",
               "Contrata menos pessoas: cada turno usa a carga semanal inteira e o posto fecha mais cedo no que "
               "sobra, sem ninguém no intervalo."),
}


def _horario_postos(postos, blocos):
    """Horário de cada linha da demanda depois do ajuste: 'Seg a Sex 07:30–16:54' (blocos na ordem das linhas)."""
    por_linha, i = [], 0
    for p in postos:
        faixas = defaultdict(list)
        for d, (ini, fim) in p["horarios"]:
            a, b = blocos[i]
            i += p["quantidade"]
            faixas[fmt_hora(a), fmt_hora(b), b - a < fim - ini].append(d)
        por_linha.append({
            "cargo": p["cargo"], "escala": p["escala"], "quantidade": p["quantidade"], "dias": p["dias"],
            "horario": "; ".join(f"{fmt_dias(dias)} {a}–{b}" for (a, b, _), dias in faixas.items()),
            "reduzido": any(r for _, _, r in faixas),
        })
    return por_linha


def _proposta_json(id_, postos, blocos, sugestoes):
    titulo, descricao = PROPOSTAS[id_]
    horas = sum(fim - ini for ini, fim in blocos) / 60
    sugestoes = _sugestoes_json(sugestoes)
    return {
        "id": id_, "titulo": titulo, "descricao": descricao,
        "salario_contratacao": None if _salario is None else round(_salario, 2),
        "postos": _horario_postos(postos, blocos),
        "horas_posto_semana": fmt_horas(horas),
        "colaboradores_equivalentes": round(horas / settings.LIMITE_HORAS_SEMANA, 1),
        "sugestoes": sugestoes,
        "por_cargo": _por_cargo([(p["cargo"], p["quantidade"], p["salario"], p["horas"]) for p in postos], sugestoes),
    }


@app.post("/dimensionar")
def dimensionar(simulacao: Simulacao):
    """
    Dimensiona a partir da demanda: cada linha diz quantos postos precisam estar ocupados nos dias e horário da
    escala (ex.: 5 postos, 'ESCALA SEGUNDA - DOMINGO (06:00 - 22:00)'). Devolve duas propostas (PROPOSTAS):
    manter o horário do posto, com o intervalo coberto, e o mínimo de colaboradores (main.minimo_de_pessoas).
    As contratações custam o salário médio dos postos, ponderado pelas horas de posto de cada linha.
    """
    global _salario
    blocos, postos, sem_horario, pesos = [], [], [], []
    for l in simulacao.linhas_com_salario():
        horarios = ler_escala(l["escala"])
        if not horarios:
            sem_horario.append({"cargo": l["cargo"], "escala": l["escala"], "quantidade": l["quantidade"]})
            continue
        novos = [(d * 1440 + ini, d * 1440 + fim) for d, (ini, fim, _) in horarios.items() for _ in range(l["quantidade"])]
        blocos += novos
        pesos.append((sum(b - a for a, b in novos), l["salario"]))
        postos.append({"cargo": l["cargo"], "escala": l["escala"], "quantidade": l["quantidade"],
                       "salario": l["salario"], "horas": sum(b - a for a, b in novos) / 60,
                       "dias": fmt_dias(list(horarios)),
                       "horarios": [(d, (d * 1440 + ini, d * 1440 + fim)) for d, (ini, fim, _) in horarios.items()]})
    with _parametros(simulacao):
        _salario = _salario_medio(pesos)
        propostas = []
        if blocos:
            propostas.append(_proposta_json("horario", postos, blocos, sugerir(sorted(blocos), com_intervalo=True)))
            ajustados, turnos, pausas = minimo_de_pessoas(blocos)
            propostas.append(_proposta_json("minimo", postos, ajustados,
                                            sugerir(sorted(ajustados), turnos=turnos, pausas=pausas)))
        return {
            "config": _config_json(),
            "sem_horario": sem_horario,
            "postos": sum(p["quantidade"] for p in postos),
            "propostas": propostas,
        }


MODELO_COLUNAS = [  # (coluna, tipo, o que é, exemplo); a ordem é a do modelo
    ("QF_CARGO", "Texto", "Cargo. O código numérico do início é removido.", "14833188 - AUXILIAR DE LIMPEZA I - DIA"),
    ("Escala", "Texto", "Dias, horário e, se houver, o intervalo: "
     "'ESCALA SEGUNDA - SEXTA (07:30 - 17:18, INTERVALO 12:00 - 13:00)' ou 'ESCALA 6 X 1 (07:30 - 16:00)'.",
     "ESCALA SEGUNDA - SEXTA (07:30 - 17:18, INTERVALO 12:00 - 13:00)"),
    ("Quantidade", "Número inteiro", "Pessoas do cargo nessa escala.", "4"),
    ("Salario", "Número (opcional)", "Salário médio mensal do cargo, para o impacto mensal e anual (x 1,75).", "1850,00"),
]
MODELO_EXEMPLOS = [
    ["200 - AUXILIAR DE LIMPEZA", "ESCALA SEGUNDA - SEXTA (07:30 - 17:18, INTERVALO 12:00 - 13:00)", 4, 1850],
    ["201 - AUXILIAR DE LIMPEZA - 6X1", "ESCALA 6 X 1 (07:30 - 16:00)", 2, 1850],
    ["202 - LIDER DE LIMPEZA", "ESCALA SEGUNDA - QUINTA (08:00 - 19:00, INTERVALO 12:00 - 13:00) / "
                               "SEXTA (08:00 - 18:00, INTERVALO 12:00 - 13:00)", 1, 2600],
    ["203 - PORTEIRO", "ESCALA SEGUNDA - SABADO (22:00 - 06:00, INTERVALO 02:00 - 03:00)", 2, 2100],
]


@app.get("/modelo")
def modelo():
    """xlsx modelo da PEC: aba PEC com o cabeçalho e linhas de exemplo, aba Instruções com o formato de cada coluna."""
    wb = Workbook()
    pec = wb.active
    pec.title = "PEC"
    pec.append([c[0] for c in MODELO_COLUNAS])
    for linha in MODELO_EXEMPLOS:
        pec.append(linha)

    inst = wb.create_sheet("Instruções")
    inst.append(["Coluna", "Tipo", "O que é", "Exemplo"])
    for coluna in MODELO_COLUNAS:
        inst.append(list(coluna))
    inst.append([])
    for texto in [
        "Cabeçalho na primeira linha da aba PEC; uma linha por cargo e escala. Outras colunas são ignoradas.",
        "Escala por dias: ESCALA <DIA> - <DIA> (HH:MM - HH:MM). Basta o início do dia (SEG, TER, QUA, QUI, SEX, SAB, DOM).",
        "Escala por ciclo: ESCALA N X M (HH:MM - HH:MM), com N + M = 7. Os N dias contam a partir de segunda.",
        "Escala dinâmica (horário diferente por dia): junte as faixas com '/', ex.: "
        "ESCALA SEGUNDA - QUINTA (08:00 - 19:00) / SEXTA (08:00 - 18:00). Um dia sozinho também vale: SABADO (08:00 - 12:00).",
        "Intervalo (opcional) dentro dos parênteses, com início e fim: (07:30 - 17:18, INTERVALO 12:00 - 13:00). "
        "Sem ele, o cálculo usa o mínimo da CLT no meio do turno. Intervalo fora do turno deixa a linha fora do cálculo.",
        "Saída menor que a entrada é turno que vira a noite, ex.: (22:00 - 06:00).",
        "Escalas fora desses formatos (12x36, vazia) somam pessoas, mas ficam fora do cálculo.",
    ]:
        inst.append([texto])

    cabecalho = Font(bold=True, color="FFFFFF")
    fundo = PatternFill("solid", fgColor="1F5EFF")
    for aba, larguras in ((pec, [40, 96, 12, 14]), (inst, [14, 16, 70, 60])):
        for celula in aba[1]:
            celula.font, celula.fill = cabecalho, fundo
        for i, largura in enumerate(larguras):
            aba.column_dimensions[chr(65 + i)].width = largura
        aba.freeze_panes = "A2"

    buf = BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="modelo_pec.xlsx"'})


@app.post("/importar")
async def importar(arquivo: UploadFile):
    """Lê um xlsx no formato da PEC e devolve as linhas por CR, para preencher o formulário."""
    if not (arquivo.filename or "").lower().endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Envie um arquivo .xlsx no formato da PEC.")
    with tempfile.TemporaryDirectory() as pasta:
        caminho = Path(pasta) / "pec.xlsx"
        caminho.write_bytes(await arquivo.read())
        try:
            por_cr = ler_linhas(caminho)
        except KeyError as coluna:
            raise HTTPException(400, f"Coluna {coluna} não encontrada. Necessárias: QF_CARGO, Escala, Quantidade.")
    campos = ("cargo", "escala", "quantidade", "salario")
    return {"centros": [{"cr": cr, "linhas": [{k: l[k] for k in campos} for l in linhas]} for cr, linhas in por_cr.items()]}


# ---------------------------------------------------------------------------
# RELATÓRIO RADAR ESCALA 5x2
# ---------------------------------------------------------------------------
@app.get("/relatorio")
def relatorio_dados():
    """Base do relatório salva no servidor (a última enviada); 404 se ainda não houver."""
    dados = relatorio.carregar()
    if dados is None:
        raise HTTPException(404, "Nenhuma base enviada ainda.")
    return dados


@app.post("/relatorio")
async def relatorio_enviar(arquivo: UploadFile):
    """Recebe a exportação de escalas e jornadas (xlsx), salva no servidor e devolve os dados do relatório."""
    if not (arquivo.filename or "").lower().endswith((".xlsx", ".xls")):
        raise HTTPException(400, "Envie a exportação em .xlsx.")
    try:
        return relatorio.salvar(await arquivo.read(), arquivo.filename)
    except Exception as erro:
        raise HTTPException(400, f"Não consegui ler a base: {erro}")


@app.post("/relatorio/cockpit")
async def relatorio_cockpit(arquivo: UploadFile):
    """Cruza uma exportação do Cockpit (csv ou xlsx) com os CRs da base: devolve as colunas por código de CR."""
    dados = relatorio.carregar()
    if dados is None:
        raise HTTPException(400, "Envie a base do relatório antes do Cockpit.")
    nome = Path(arquivo.filename or "cockpit.csv").name
    with tempfile.TemporaryDirectory() as pasta:
        caminho = Path(pasta) / nome
        caminho.write_bytes(await arquivo.read())
        try:
            return relatorio.cockpit(caminho, nome, [c[0] for c in dados["crs"]])
        except Exception as erro:
            raise HTTPException(400, f"Não consegui ler o Cockpit: {erro}")
